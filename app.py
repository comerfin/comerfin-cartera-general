import os
import base64
import html
import io
import subprocess
import sys
import datetime
from zoneinfo import ZoneInfo
import streamlit as st
import pandas as pd
import numpy as np

from data_loader import cargar_datos_reporte, obtener_comercios_propios, procesar_excel_a_df
from metrics import calcular_metricas, calcular_proyeccion, calcular_dias_laborables, normalizar_estado
from analytics import calcular_semaforo, evaluar_cumplimiento_objetivo, calcular_tendencia_mes_a_mes

# El acceso compartido utiliza otro punto de entrada; la URL nunca habilita la cartera propia.
solo_general = True
nombre_visitante = globals().get("NOMBRE_VISITANTE", "") if solo_general else ""

def ruta_logo(nombre):
    carpeta = os.path.dirname(__file__)
    local = os.path.join(carpeta, "assets", nombre)
    return local if os.path.exists(local) else os.path.join(carpeta, nombre)

st.set_page_config(
    page_title="Dashboard Cartera Comercios | Comercial Group",
    page_icon=ruta_logo("comercial-group.png"),
    layout="wide",
    initial_sidebar_state="auto"
)

st.set_option("client.disableDataExport", True)

# Estilos visuales modernos
st.markdown("""
<style>
    .metric-card {
        box-sizing: border-box;
        height: 180px;
        display: flex;
        flex-direction: column;
        background-color: #ffffff;
        border-radius: 10px;
        padding: 16px 20px;
        box-shadow: 0 2px 8px rgba(0,0,0,0.06);
        border: 1px solid #eef0f3;
        margin-bottom: 15px;
    }
    .conversion-card {
        box-sizing: border-box;
        height: 180px;
        display: flex;
        flex-direction: column;
        border-top: 4px solid #f28c28;
        box-shadow: 0 4px 16px rgba(0,0,0,0.10);
    }
    .metric-card .metric-title {min-height: 40px;}
    .metric-card .metric-value {line-height: 1.2;}
    .metric-card .metric-sub {margin-top: auto; line-height: 1.35;}
    .conversion-card .metric-value {font-size: 2rem; line-height: 1.2;}
    .conversion-card .metric-sub {margin-top: auto; line-height: 1.35;}
    .metric-title {
        font-size: 0.85rem;
        color: #64748b;
        font-weight: 600;
        text-transform: uppercase;
        margin-bottom: 6px;
    }
    .metric-value {
        font-size: 1.8rem;
        font-weight: 800;
        color: #1e293b;
    }
    .metric-sub {
        font-size: 0.85rem;
        color: #059669;
        font-weight: 600;
    }
    [data-testid="stMainBlockContainer"] {max-width: 1600px;}
    [data-testid="stColumn"] {min-width: 0;}
    .metric-value {white-space: nowrap; font-size: clamp(1.15rem, 1.85vw, 1.8rem);}
    .conversion-card .metric-value {font-size: clamp(1.35rem, 2vw, 2rem);}
    [data-testid="stMetricValue"] {font-size: clamp(1.25rem, 2.1vw, 2.25rem);}
    [data-testid="stBaseButton-secondary"] {min-height: 44px;}
    @media (min-width: 641px) and (max-width: 1100px) {
        [data-testid="stHorizontalBlock"] {flex-wrap: wrap; gap: 1rem;}
        [data-testid="stHorizontalBlock"] > [data-testid="stColumn"] {flex: 1 1 calc(50% - 1rem) !important; min-width: 0 !important;}
        .metric-value {font-size: clamp(1.25rem, 2.7vw, 1.8rem);}
    }
    @media (max-width: 640px) {
        [data-testid="stMainBlockContainer"] {padding: 3.5rem 1rem 2rem;}
        [data-testid="stHorizontalBlock"] {flex-wrap: wrap; gap: .75rem;}
        [data-testid="stHorizontalBlock"] > [data-testid="stColumn"] {flex: 1 1 100% !important; width: 100% !important; min-width: 0 !important;}
        .metric-card {padding: 16px; margin-bottom: 0;}
        .metric-value {font-size: clamp(1.2rem, 6vw, 1.8rem);}
        .metric-card {height: 180px;}
        .conversion-card .metric-value {font-size: 1.8rem;}
        [data-testid="stMetricValue"] {font-size: clamp(1.2rem, 6vw, 2rem);}
        [role="dialog"] {max-width: calc(100vw - 16px) !important; margin: 8px !important;}
        [data-testid="stDataFrame"] {max-width: 100%;}
    }
    .st-key-venta_acumulada_dia {
        background: linear-gradient(110deg, #ffffff 0%, #fff7ed 100%);
        border: 1px solid #fed7aa !important;
        border-left: 6px solid #f28c28 !important;
        border-radius: 14px;
        padding: 22px !important;
        box-shadow: 0 6px 22px rgba(242,140,40,.13);
        margin-bottom: 14px;
        color: #1e293b;
    }
    .st-key-venta_acumulada_dia h3, .st-key-venta_acumulada_dia p,
    .st-key-venta_acumulada_dia [data-testid="stMetricValue"],
    .st-key-venta_acumulada_dia [data-testid="stMetricLabel"] {color: #1e293b;}
    .st-key-venta_acumulada_dia [data-testid="stMetricValue"] {font-size: clamp(1.7rem, 3vw, 2.8rem); font-weight: 800;}
    .st-key-resumen_consultas_periodos [data-testid="stMetricValue"] {font-size: clamp(1.2rem, 1.8vw, 1.65rem);}
    [data-testid="stSidebar"] a[href="https://forms.gle/1u6eQK39Hh7xZGjn7"] {
        background: #ffddb3; color: #000000; border: 1px solid #f5bb78;
        border-radius: 10px; font-weight: 700;
    }
    [data-testid="stSidebar"] a[href="https://forms.gle/1u6eQK39Hh7xZGjn7"]:hover {
        background: #ffd09a; color: #000000; border-color: #eda451;
    }
    [data-testid="stSidebar"] a[href="https://forms.gle/1u6eQK39Hh7xZGjn7"] p {
        color: #000000; font-weight: 700;
    }
    [data-testid="stSidebar"] a[href="https://forms.gle/1u6eQK39Hh7xZGjn7"] [data-testid="stIconMaterial"] {
        color: #000000; font-size: 1.4rem;
    }
    .daily-counts {display:flex; flex-wrap:wrap; gap:8px; margin:4px 0 8px;}
    .daily-counts span {background:#fff; border:1px solid #fed7aa; border-radius:20px; padding:6px 12px; color:#7c2d12; font-size:.85rem;}
    .badge-verde {
        background-color: #d1fae5;
        color: #065f46;
        padding: 4px 10px;
        border-radius: 6px;
        font-weight: 700;
        font-size: 0.85rem;
    }
    .badge-amarillo {
        background-color: #fef3c7;
        color: #92400e;
        padding: 4px 10px;
        border-radius: 6px;
        font-weight: 700;
        font-size: 0.85rem;
    }
    .badge-rojo {
        background-color: #fee2e2;
        color: #991b1b;
        padding: 4px 10px;
        border-radius: 6px;
        font-weight: 700;
        font-size: 0.85rem;
    }
</style>
""", unsafe_allow_html=True)


def pesos(valor):
    if pd.isna(valor):
        return "—"
    return "$" + format(float(valor), ",.0f").replace(",", ".") + ".-"


def archivo_excel(datos):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    libro = Workbook()
    hoja = libro.active
    hoja.title = "Reporte"
    hoja.append(list(datos.columns))
    for fila in datos.itertuples(index=False, name=None):
        valores = []
        for valor in fila:
            if pd.isna(valor):
                valor = None
            elif hasattr(valor, "item"):
                valor = valor.item()
            valores.append(valor)
        hoja.append(valores)
        # Los textos del portal se exportan como texto, nunca como fórmulas.
        for celda in hoja[hoja.max_row]:
            if isinstance(celda.value, str):
                celda.data_type = "s"
    for celda in hoja[1]:
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor="243246")
        celda.alignment = Alignment(wrap_text=True)
    for indice, nombre in enumerate(datos.columns, 1):
        hoja.column_dimensions[get_column_letter(indice)].width = min(44, max(18, len(str(nombre)) + 4))
        monetaria = nombre == "Monto" or "($)" in str(nombre)
        for fila in hoja.iter_rows(min_row=2, min_col=indice, max_col=indice):
            if monetaria:
                fila[0].number_format = '"$"#,##0".-"'
            elif "(%)" in str(nombre):
                fila[0].number_format = '0.00"%"'
    hoja.freeze_panes = "A2"
    hoja.auto_filter.ref = hoja.dimensions
    salida = io.BytesIO()
    libro.save(salida)
    return salida.getvalue()


def descargar_excel(datos, nombre, clave):
    st.download_button("📥 Descargar Excel", archivo_excel(datos), file_name=nombre + ".xlsx",
                       mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", key=clave)


def ticket_promedio(monto, cantidad):
    return monto / cantidad if cantidad else None


@st.cache_data(ttl=60)
def get_data(report_version, cartera_version=None):
    if report_version is None:
        return pd.DataFrame()
    report_path = os.path.join(os.path.dirname(__file__), "downloads", "Reporte.xlsx")
    datos = procesar_excel_a_df(report_path)
    if "Nombre" not in datos.columns:
        datos["Nombre"] = pd.read_excel(report_path, usecols=["Nombre"])["Nombre"]
    return datos


@st.cache_data(ttl=60)
def cargar_historial_liquidaciones(version):
    if version is None:
        return pd.DataFrame()
    from data_loader import limpiar_monto
    ruta = os.path.join(os.path.dirname(__file__), "downloads", "HistorialLiquidaciones.xlsx")
    datos = pd.read_excel(ruta, usecols=["Solicitud", "Comercio", "Estado", "FechaLiquidacion", "Capital", "MontoMaximo"])
    capital = datos["Capital"].apply(limpiar_monto)
    maximo = datos["MontoMaximo"].apply(limpiar_monto)
    datos["Monto"] = capital.where(capital > 0, maximo)
    datos["Comercio"] = datos["Comercio"].astype(str).str.strip()
    return datos


def actualizar_desde_portal():
    base = os.path.dirname(os.path.abspath(__file__))
    entorno = os.environ.copy()
    try:
        for nombre in ("PORTAL_URL", "PORTAL_USER", "PORTAL_PASSWORD"):
            if nombre in st.secrets:
                entorno[nombre] = str(st.secrets[nombre])
    except FileNotFoundError:
        pass
    try:
        result = subprocess.run(
            [sys.executable, "-u", os.path.join(base, "actualizar_reporte.py")],
            cwd=base, capture_output=True, text=True, timeout=180, env=entorno,
        )
    except subprocess.TimeoutExpired:
        st.error("Prendo no respondió en 3 minutos. El reporte anterior se conserva. Volvé a intentar.")
        return False
    if result.returncode != 0:
        st.error("No se pudo actualizar el reporte. Se conserva la información anterior.")
        with st.expander("Detalle de la actualización", expanded=True):
            st.text(result.stdout or "El actualizador no pudo iniciarse. Revisá el entorno de Python del proyecto.")
        return False
    st.cache_data.clear()
    st.session_state["reporte_actualizado"] = True
    return True


# Cargar datos
with st.spinner("Cargando reporte de operaciones de Prendo / MG Group..."):
    df_raw = get_data(os.stat(os.path.join(os.path.dirname(__file__), "downloads", "Reporte.xlsx")).st_mtime_ns if os.path.exists(os.path.join(os.path.dirname(__file__), "downloads", "Reporte.xlsx")) else None, os.stat(os.path.join(os.path.dirname(__file__), "cartera.json")).st_mtime_ns)

if df_raw.empty:
    st.error("No se encontró el archivo de reporte en `downloads/Reporte.xlsx`. Haz clic en 'Descargar Datos' para iniciar la extracción automática.")
    if st.button("🚀 Iniciar Descarga del Portal"):
        with st.spinner("Descargando reporte de Prendo..."):
            actualizado = actualizar_desde_portal()
        if actualizado:
            st.rerun()
    st.stop()

if solo_general:
    df_raw = df_raw.drop(columns=["Es_Mi_Cartera"], errors="ignore").copy()

import json

CONFIG_TEXTOS_PATH = os.path.join(os.path.dirname(__file__), "config_textos.json")

def cargar_textos() -> dict:
    if os.path.exists(CONFIG_TEXTOS_PATH):
        try:
            with open(CONFIG_TEXTOS_PATH, "r", encoding="utf-8") as f:
                return json.load(f)
        except Exception:
            pass
    return {}

txt = cargar_textos()

# SIDEBAR: Filtros y Configuración
def logo_data(nombre):
    with open(ruta_logo(nombre), "rb") as archivo:
        return "data:image/png;base64," + base64.b64encode(archivo.read()).decode("ascii")


st.markdown("""<style>
.sidebar-brand {padding: 4px 0 8px; display:flex; flex-direction:column; align-items:center; gap:18px;}
.sidebar-brand .comercial-logo {width:165px; max-width:100%; border-radius:16px; display:block;}
.sidebar-brand .powered-copy {font-size:13px; line-height:1.5; text-align:center; margin:0; opacity:.82;}
.sidebar-brand .comerfin-logo {width:140px; max-width:85%; border-radius:8px; display:block; margin-top:-8px;}
.dashboard-brand {display:flex; align-items:center; flex-wrap:wrap; gap:24px; margin:0 0 18px;}
.dashboard-brand img {width:300px; max-width:100%; height:auto; border-radius:14px; display:block;}
.dashboard-brand h1 {font-size:clamp(25px,3vw,38px); line-height:1.18; letter-spacing:-.035em; margin:0; padding:0; font-weight:750;}
@media(max-width:600px) {.dashboard-brand {gap:16px;} .dashboard-brand img {width:270px;}}
</style>""", unsafe_allow_html=True)
st.sidebar.markdown(f'<div class="sidebar-brand"><img class="comercial-logo" src="{logo_data("comercial-group.png")}" alt="Comercial Group"></div>', unsafe_allow_html=True)
st.sidebar.divider()

st.sidebar.link_button(
    "Dar de alta un comercio",
    "https://forms.gle/1u6eQK39Hh7xZGjn7",
    icon=":material/add_business:",
    type="primary",
    use_container_width=True,
    help="Abrir el formulario de alta de comercios en una nueva pestaña.",
)
st.sidebar.caption("Ingresá una nueva solicitud de adhesión.")
st.sidebar.divider()

# Selector Principal de Vista
seccion = st.sidebar.radio("Sección", ["Dashboard", "Reportes y gráficos"], key="seccion_principal")
if solo_general:
    vista_seleccionada = "🌐 Cartera General (Red Completa)"
else:
    vista_seleccionada = st.sidebar.radio(
        "Seleccionar Vista:",
        ["⭐ Mi Cartera (Mis Comercios)", "🌐 Cartera General (Red Completa)"],
        index=0
    )

# Filtro de Cartera
if "Mi Cartera" in vista_seleccionada:
    df_base = df_raw[df_raw["Es_Mi_Cartera"] == True].copy()
    st.sidebar.success(f"🎯 Foco: Cartera Propia ({len(df_base):,} consultas)")
else:
    df_base = df_raw.copy()
    st.sidebar.info(f"🌐 Red General: {len(df_base):,} consultas")

# Filtro de Fechas con formato DD/MM/YYYY
st.sidebar.subheader("📅 Rango de Fechas")
fechas_validas = df_base["FechaAlta"].dropna()
min_date = fechas_validas.min().date() if not fechas_validas.empty else datetime.date(2026, 10, 1)
max_date = fechas_validas.max().date() if not fechas_validas.empty else datetime.date(2026, 10, 14)

fecha_rango = st.sidebar.date_input(
    "Filtrar período:",
    value=(max(min_date, max_date.replace(day=1)), max_date),
    min_value=min_date,
    max_value=max_date,
    format="DD/MM/YYYY"
)


if isinstance(fecha_rango, tuple) and len(fecha_rango) == 2:
    f_desde, f_hasta = fecha_rango
    df_filtrado = df_base[(df_base["FechaAlta"].dt.date >= f_desde) & (df_base["FechaAlta"].dt.date <= f_hasta)].copy()
else:
    f_desde, f_hasta = min_date, max_date
    df_filtrado = df_base.copy()

# Filtro de Comercio / Cadena
st.sidebar.subheader("🏪 Filtro por Comercio")
comercios_disponibles = sorted([c for c in df_filtrado["Comercio"].unique() if str(c).strip()])
comercios_sel = st.sidebar.multiselect(
    "Seleccionar Comercios:",
    options=comercios_disponibles,
    default=[]
)
if comercios_sel:
    df_filtrado = df_filtrado[df_filtrado["Comercio"].isin(comercios_sel)]

# Filtro por Localidad / Zona
st.sidebar.subheader("📍 Filtro por Localidad")
localidades_disponibles = sorted([loc for loc in df_filtrado["Localidad"].unique() if str(loc).strip()])
localidades_sel = st.sidebar.multiselect(
    "Seleccionar Localidad / Zona:",
    options=localidades_disponibles,
    default=[]
)
if localidades_sel:
    df_filtrado = df_filtrado[df_filtrado["Localidad"].isin(localidades_sel)]

# Botón para refrescar

st.sidebar.divider()
if st.session_state.pop("reporte_actualizado", False):
    st.sidebar.success("¡Reporte descargado y datos actualizados!")
if st.sidebar.button("🔄 Actualizar Datos del Portal"):
    with st.sidebar.status("Conectando a Prendo y descargando reporte (máximo 3 minutos)...") as status:
        actualizado = actualizar_desde_portal()
        status.update(label="Actualización completada" if actualizado else "No se pudo actualizar", state="complete" if actualizado else "error")
    if actualizado:
        st.rerun()

report_path = os.path.join(os.path.dirname(__file__), "downloads", "Reporte.xlsx")
if os.path.exists(report_path):
    from zoneinfo import ZoneInfo
    ultima_actualizacion = datetime.datetime.fromtimestamp(os.path.getmtime(report_path), ZoneInfo("America/Argentina/Buenos_Aires"))
    st.sidebar.info("Última actualización del reporte: " + ultima_actualizacion.strftime("%d/%m/%Y %H:%M:%S") + " (Argentina)")

st.sidebar.markdown(f'<div style="margin-top:32px; padding:16px 0 8px; text-align:left; border-top:1px solid #e5e7eb;"><img src="{logo_data("comerfin-powered-data.png")}" alt="ComerFin — Powered &amp; Data Analitycs" style="width:110px; max-width:55%; height:auto;"><div style="font-size:10px; opacity:.65; margin-top:6px;">Comercial Intelligent</div></div>', unsafe_allow_html=True)

# ----------------- PANEL PRINCIPAL -----------------
main_title = txt.get("titulo_mi_cartera", "Dashboard Cartera Comercios") if "Mi Cartera" in vista_seleccionada else txt.get("titulo_cartera_general", "🌐 Dashboard: Red General de Mercado")
st.markdown(f'<div class="dashboard-brand"><img src="{logo_data("prendo.png")}" alt="Prendo — Rápido. Seguro. Fácil."><h1>{html.escape(main_title)}</h1></div>', unsafe_allow_html=True)
st.caption(f"Período: **{f_desde.strftime('%d/%m/%Y')} al {f_hasta.strftime('%d/%m/%Y')}** | Operaciones analizadas: **{len(df_filtrado):,}**")

if seccion == "Reportes y gráficos":
    from reportes import mostrar_reportes
    mostrar_reportes(df_filtrado, pesos)
    st.stop()

# Transformar a registros para el motor de métricas
records = []
for _, row in df_filtrado.iterrows():
    records.append({
        "estado": row["Estado"],
        "monto": row["Monto"],
        "comercio": row["Comercio"],
        "cadena": row["Cadena"],
    })

m = calcular_metricas(records)

# Calcular métricas globales de toda la cartera propia (para referencia fija)
if not solo_general:
    df_mi_cartera_full = df_raw[df_raw["Es_Mi_Cartera"] == True].copy()
    rec_cartera_full = [{"estado": r["Estado"], "monto": r["Monto"]} for _, r in df_mi_cartera_full.iterrows()]
    m_cartera_full = calcular_metricas(rec_cartera_full)

@st.dialog("Listado de operaciones", width="large")
def mostrar_operaciones(tipo, datos):
    estados = datos["Estado"].map(normalizar_estado)
    estados_por_tipo = {
        "liquidado": ["LIQUIDADA"],
        "analisis": ["EN_ANALISIS"],
        "pipeline": ["LIQUIDADA", "EN_ANALISIS", "APROBADA"],
    }
    permitidos = estados_por_tipo[tipo]
    operaciones = datos[estados.isin(permitidos)].copy()
    titulo = {"liquidado": "Total liquidado", "analisis": "Préstamos en análisis", "pipeline": "Venta total (Pipeline)"}[tipo]
    st.subheader(titulo)
    st.caption(f"Período: {f_desde.strftime('%d/%m/%Y')} al {f_hasta.strftime('%d/%m/%Y')}. Se aplican los filtros de cartera, comercio y localidad.")
    st.write(f"**{len(operaciones)} operaciones · {pesos(operaciones['Monto'].sum())}**")
    promedio = ticket_promedio(operaciones["Monto"].sum(), len(operaciones))
    st.metric("Ticket promedio", pesos(promedio) if promedio is not None else "—")
    if operaciones.empty:
        st.info("No hay operaciones para los filtros seleccionados.")
        return
    columnas = ["Nombre", "Comercio", "Estado", "FechaAlta", "FechaAprobacion", "FechaLiquidacion", "Monto", "Solicitud"]
    listado = operaciones[[c for c in columnas if c in operaciones.columns]].copy()
    listado = listado.sort_values("FechaAlta", ascending=False)
    for columna in ["FechaAlta", "FechaAprobacion", "FechaLiquidacion"]:
        if columna in listado:
            listado[columna] = pd.to_datetime(listado[columna], dayfirst=True, errors="coerce").dt.strftime("%d/%m/%Y").fillna("—")
    listado = listado.rename(columns={"Nombre": "Nombre y apellido", "FechaAlta": "Fecha de operación", "FechaAprobacion": "Fecha de aprobación", "FechaLiquidacion": "Fecha de liquidación", "Solicitud": "Solicitud Nº"})
    if "Nombre y apellido" not in listado:
        listado.insert(0, "Nombre y apellido", "No informado")
    listado["Nombre y apellido"] = listado["Nombre y apellido"].fillna("No informado")
    st.dataframe(listado.style.format({"Monto": pesos}), hide_index=True, use_container_width=True, height=430)
    st.markdown(f"**Total del monto colocado: {pesos(operaciones['Monto'].sum())}**")
    st.caption("Fecha de operación corresponde al alta de la solicitud. También se muestran aprobación y liquidación.")
    descargar_excel(listado, "operaciones_" + tipo, "excel_" + tipo)


def detalle_creditos_diarios(datos, clave):
    listado = datos[[c for c in ["Solicitud", "Nombre", "FechaAprobacion", "Monto", "Comercio", "Estado"] if c in datos.columns]].copy()
    if "Nombre" not in listado:
        listado["Nombre"] = "No informado"
    listado["Nombre"] = listado["Nombre"].fillna("No informado")
    listado["FechaAprobacion"] = pd.to_datetime(listado["FechaAprobacion"], dayfirst=True, errors="coerce").dt.strftime("%d/%m/%Y").fillna("—")
    listado = listado.rename(columns={"Solicitud": "Solicitud Nº", "Nombre": "Nombre y apellido", "FechaAprobacion": "Fecha de aprobación"})
    if listado.empty:
        st.info("No hay créditos en este estado para la fecha consultada.")
    else:
        st.dataframe(listado.style.format({"Monto": pesos}), hide_index=True, use_container_width=True)
    st.markdown(f"**Total del monto colocado: {pesos(datos['Monto'].sum())}**")
    descargar_excel(listado, "colocacion_" + clave, "excel_diario_" + clave)


@st.dialog("Créditos de la fecha", width="large")
def mostrar_creditos_diarios(datos, estado, fecha):
    etiquetas = {"todos": "Todas las operaciones", "APROBADA": "Aprobados", "EN_ANALISIS": "En análisis", "LIQUIDADA": "Liquidados"}
    seleccion = datos if estado == "todos" else datos[datos["Estado"].map(normalizar_estado) == estado]
    st.subheader(etiquetas[estado])
    st.caption(f"Fecha de aprobación: {fecha.strftime('%d/%m/%Y')} · {len(seleccion)} créditos. Se aplican los filtros de cartera, comercio y localidad.")
    detalle_creditos_diarios(seleccion, estado + "_" + fecha.isoformat())


# Consulta diaria independiente del rango general; hoy es la vista inicial.
ahora_argentina = datetime.datetime.now(ZoneInfo("America/Argentina/Buenos_Aires"))
hoy_argentina = ahora_argentina.date()
base_diaria = df_base
if comercios_sel:
    base_diaria = base_diaria[base_diaria["Comercio"].isin(comercios_sel)]
if localidades_sel:
    base_diaria = base_diaria[base_diaria["Localidad"].isin(localidades_sel)]
with st.container(border=True, key="venta_acumulada_dia"):
    if nombre_visitante:
        st.caption(f"Bienvenido, {nombre_visitante}")
    st.subheader("⚡ Funnel · Acumulado mensual y parcial diario")
    modo_diario = st.radio("Consultar día", ["Hoy", "Elegir fecha"], horizontal=True, key="modo_consulta_diaria")
    fecha_consulta = hoy_argentina
    if modo_diario == "Elegir fecha":
        fechas_reporte = df_raw["FechaAlta"].dropna()
        inicio_reporte = min(hoy_argentina, fechas_reporte.min().date()) if not fechas_reporte.empty else hoy_argentina
        fecha_consulta = st.date_input("Fecha de consulta", value=hoy_argentina, min_value=inicio_reporte, max_value=hoy_argentina, format="DD/MM/YYYY", key="fecha_consulta_diaria")
    st.caption(("Fecha actual: " if fecha_consulta == hoy_argentina else "Fecha consultada: ") + fecha_consulta.strftime("%d/%m/%Y"))
    consultas_dia = base_diaria[
        pd.to_datetime(base_diaria["FechaAlta"], dayfirst=True, errors="coerce").dt.date == fecha_consulta
    ]
    metricas_consultas_dia = calcular_metricas([
        {"estado": estado} for estado in consultas_dia["Estado"]
    ])
    fechas_alta_periodo = pd.to_datetime(base_diaria["FechaAlta"], dayfirst=True, errors="coerce").dt.date
    inicio_mes_consulta = fecha_consulta.replace(day=1)
    consultas_mes = base_diaria[(fechas_alta_periodo >= inicio_mes_consulta) & (fechas_alta_periodo <= fecha_consulta)]
    metricas_consultas_mes = calcular_metricas([{"estado": estado} for estado in consultas_mes["Estado"]])
    porcentaje_oferta = metricas_consultas_dia["pct_aprobacion"]
    semaforo_oferta = calcular_semaforo(porcentaje_oferta, "pct_aprobacion")
    with st.container(key="resumen_consultas_periodos"):
        st.markdown("**Acumulado del mes**")
        mensual_consultas, mensual_oferta = st.columns(2)
        mensual_consultas.metric("Consultas brutas acumuladas del mes", f"{metricas_consultas_mes['total_consultas']:,}")
        mensual_oferta.metric("Approval rate del mes", f"{metricas_consultas_mes['pct_aprobacion']:.2f}%" if metricas_consultas_mes["total_consultas"] else "—")
        st.caption(f"{inicio_mes_consulta:%d/%m/%Y} al {fecha_consulta:%d/%m/%Y} · {metricas_consultas_mes['total_con_oferta']:,} consultas con oferta de {metricas_consultas_mes['total_consultas']:,} recibidas.")
        st.markdown("**Parcial de hoy**" if fecha_consulta == hoy_argentina else "**Parcial del día consultado**")
        consultas_col, oferta_col = st.columns(2)
        consultas_col.metric("Consultas brutas acumuladas del día", f"{metricas_consultas_dia['total_consultas']:,}")
        oferta_col.metric("Approval rate del día", f"{porcentaje_oferta:.2f}%" if metricas_consultas_dia["total_consultas"] else "—")
        st.caption(f"{fecha_consulta:%d/%m/%Y} · {metricas_consultas_dia['total_con_oferta']:,} consultas con oferta de {metricas_consultas_dia['total_consultas']:,} recibidas.")
    estados_diarios = base_diaria["Estado"].map(normalizar_estado)
    ventas_dia = base_diaria[
        (pd.to_datetime(base_diaria["FechaAprobacion"], dayfirst=True, errors="coerce").dt.date == fecha_consulta)
        & estados_diarios.isin(["APROBADA", "EN_ANALISIS", "LIQUIDADA"])
    ]
    st.metric("Monto acumulado hoy" if fecha_consulta == hoy_argentina else "Monto colocado en la fecha", pesos(ventas_dia["Monto"].sum()))
    conteos_dia = ventas_dia["Estado"].map(normalizar_estado).value_counts()
    for columna, estado, etiqueta in zip(st.columns(4), ["todos", "APROBADA", "EN_ANALISIS", "LIQUIDADA"], ["operaciones", "aprobados", "en análisis", "liquidados"]):
        cantidad = len(ventas_dia) if estado == "todos" else int(conteos_dia.get(estado, 0))
        if columna.button(f"{cantidad} {etiqueta}", key="diario_" + estado, use_container_width=True):
            mostrar_creditos_diarios(ventas_dia, estado, fecha_consulta)
    with st.expander("Más info", expanded=False):
        st.caption("Oferta = incompletas con oferta + en análisis + aprobadas + liquidadas, sobre las consultas dadas de alta ese día y según su estado actual. Sin consultas, el porcentaje no aplica.")
        st.caption("Según fecha de aprobación y estado actual del crédito. Respeta cartera, comercios y localidad; es independiente del rango general. Consulta limitada al período incluido en el reporte descargado.")
    if not ventas_dia.empty:
        with st.expander("Ver operaciones de la fecha"):
            detalle_creditos_diarios(ventas_dia, "todos_expandido_" + fecha_consulta.isoformat())
    reporte_actualizado = datetime.datetime.fromtimestamp(os.path.getmtime(report_path), ZoneInfo("America/Argentina/Buenos_Aires"))
    st.caption("Datos descargados el " + reporte_actualizado.strftime("%d/%m/%Y · %H:%M") + " hs · Argentina")
    if reporte_actualizado.date() < fecha_consulta:
        st.info("El reporte disponible es anterior a la fecha consultada. Usá Actualizar Datos del Portal.")

# RENDER DE KPIS PRINCIPALES
c1, c2, c3, c4 = st.columns(4)

with c1:
    st.markdown(f"""
    <div class="metric-card">
        <div class="metric-title">{txt.get('kpi_consultas_titulo', '📋 Consultas Totales')}</div>
        <div class="metric-value">{m['total_consultas']:,}</div>
        <div class="metric-sub">{txt.get('kpi_consultas_sub', 'Base total recibida')}</div>
    </div>
    """, unsafe_allow_html=True)

with c2:
    sem_aprob = calcular_semaforo(m['pct_aprobacion'], "pct_aprobacion")
    st.markdown(f"""
    <div class="metric-card">
        <div class="metric-title">{txt.get('kpi_aprobacion_titulo', '🎯 % Aprobación (Oferta)')}</div>
        <div class="metric-value">{m['pct_aprobacion']}% <span style="font-size:1.2rem;">{sem_aprob['icono']}</span></div>
        <div class="metric-sub">{m['total_con_oferta']:,} {txt.get('kpi_aprobacion_sub', 'tuvieron oferta aprobada')}</div>
    </div>
    """, unsafe_allow_html=True)

with c3:
    st.markdown(f"""
    <div class="metric-card">
        <div class="metric-title">{txt.get('kpi_venta_total_titulo', '💼 Venta Total (Pipeline)')}</div>
        <div class="metric-value">{pesos(m['monto_venta_total'])}</div>
        <div class="metric-sub">En análisis + aprobados + liquidados</div>
    </div>
    """, unsafe_allow_html=True)

    if st.button(f"{m['venta_total_cant']} operaciones (Análisis + Aprobados + Liquidados)", key="abrir_pipeline", use_container_width=True):
        mostrar_operaciones("pipeline", df_filtrado)

with c4:
    st.markdown(f"""
    <div class="metric-card">
        <div class="metric-title">{txt.get('kpi_liquidado_titulo', '💰 Total Colocado (Liquidado)')}</div>
        <div class="metric-value">{pesos(m['monto_liquidado'])}</div>
        <div class="metric-sub">Créditos efectivamente desembolsados</div>
    </div>
    """, unsafe_allow_html=True)

    if st.button(f"{m['total_liquidado_cant']} créditos desembolsados", key="abrir_liquidados", use_container_width=True):
        mostrar_operaciones("liquidado", df_filtrado)


st.subheader("Ticket promedio de créditos")
with st.expander("Más info", expanded=False):
    st.caption("Monto total dividido por la cantidad de créditos en cada estado. El conjunto incluye aprobados, en análisis y liquidados, respetando los filtros seleccionados.")
tickets = [
    ("Tiket promedio total", m["monto_venta_total"], m["venta_total_cant"]),
    ("Ticket promedio aprobados", m["monto_aprobadas"], m["aprobadas"]),
    ("Ticket promedio en análisis", m["monto_en_analisis"], m["en_analisis"]),
    ("Ticket promedio liquidados", m["monto_liquidado"], m["liquidadas"]),
]
for columna, (titulo, monto, cantidad) in zip(st.columns(4), tickets):
    promedio = ticket_promedio(monto, cantidad)
    columna.metric(titulo, pesos(promedio) if promedio is not None else "—")
    columna.caption(f"{cantidad} créditos" if cantidad else "Sin créditos en este estado")

# SEGUNDA FILA: RATIOS DE CONVERSIÓN Y EMBUDO
st.subheader("Conversión y cierre")
k1, k2, k3, k4 = st.columns(4)

with k1:
    sem_conv_b = calcular_semaforo(m['pct_conversion_bruta'], "pct_conversion_bruta")
    st.markdown(f"""
    <div class="metric-card conversion-card">
        <div class="metric-title">{txt.get('kpi_conversion_bruta_titulo', '⚡ Conversión Bruta (Liquidado)')}</div>
        <div class="metric-value">{m['pct_conversion_bruta']}% {sem_conv_b['icono']}</div>
        <div class="metric-sub">{txt.get('kpi_conversion_bruta_sub', 'Liquidadas / Consultas')}</div>
    </div>
    """, unsafe_allow_html=True)

with k2:
    sem_conv_tot = calcular_semaforo(m['pct_conversion_bruta_total'], "pct_conversion_bruta")
    st.markdown(f"""
    <div class="metric-card conversion-card">
        <div class="metric-title">{txt.get('kpi_conversion_vta_total_titulo', '⚡ Conversión Bruta (Vta Total)')}</div>
        <div class="metric-value">{m['pct_conversion_bruta_total']}% {sem_conv_tot['icono']}</div>
        <div class="metric-sub">{txt.get('kpi_conversion_vta_total_sub', 'Venta Total / Consultas')}</div>
    </div>
    """, unsafe_allow_html=True)

with k3:
    st.markdown(f"""
    <div class="metric-card conversion-card">
        <div class="metric-title">{txt.get('kpi_conversion_oferta_titulo', '🤝 Conversión s/Oferta')}</div>
        <div class="metric-value">{m['pct_conversion_sobre_oferta']}%</div>
        <div class="metric-sub">{txt.get('kpi_conversion_oferta_sub', 'Efectividad de cierre de ofertas')}</div>
    </div>
    """, unsafe_allow_html=True)

with k4:
    st.markdown(f"""
    <div class="metric-card conversion-card">
        <div class="metric-title">{txt.get('kpi_en_analisis_titulo', '⏳ En análisis')}</div>
        <div class="metric-value">{m['en_analisis']} <span style="font-size:1rem;color:#f59e0b;">({pesos(m['monto_en_analisis'])})</span></div>
        <div class="metric-sub">{txt.get('kpi_en_analisis_sub', 'Pueden estar pendientes de remito o factura, o esperando el cambio de estado a aprobado.')}</div>
    </div>
    """, unsafe_allow_html=True)

    if st.button(f"Ver {m['en_analisis']} préstamos en análisis", key="abrir_en_analisis", use_container_width=True):
        mostrar_operaciones("analisis", df_filtrado)

st.divider()

# PROYECCIÓN DE CIERRE DE MES (RUN-RATE)
st.subheader(txt.get("seccion_proyeccion_titulo", "📈 Proyección de Cierre de Mes (Run-Rate)"))


# Cálculo de días hábiles del mes
hoy = datetime.date.today()
primer_dia_mes = hoy.replace(day=1)
# Último día del mes
if hoy.month == 12:
    ultimo_dia_mes = datetime.date(hoy.year, 12, 31)
else:
    ultimo_dia_mes = datetime.date(hoy.year, hoy.month + 1, 1) - datetime.timedelta(days=1)

dias_trabajados = calcular_dias_laborables(primer_dia_mes, hoy, incluir_sabados=True)
dias_restantes = calcular_dias_laborables(hoy + datetime.timedelta(days=1), ultimo_dia_mes, incluir_sabados=True)

proy = calcular_proyeccion(
    monto_acumulado=m["monto_venta_total"],
    cant_acumulada=m["venta_total_cant"],
    dias_trabajados=max(1, dias_trabajados),
    dias_restantes=dias_restantes
)

p1, p2, p3, p4 = st.columns(4)
with p1:
    st.metric("Días Hábiles Transcurridos", f"{proy['dias_trabajados']} días")
with p2:
    st.metric("Días Hábiles Restantes", f"{proy['dias_restantes']} días")
with p3:
    st.metric("Ritmo Diario Actual", f"{pesos(proy['promedio_diario_monto'])} / día", f"{proy['promedio_diario_cant']:.1f} ventas/día")
with p4:
    st.metric("🎯 CIERRE TOTAL ESTIMADO", f"{pesos(proy['proyeccion_total_monto'])}", f"{proy['proyeccion_total_cant']} ventas proyectadas")

st.divider()

# DETALLE Y RANKING POR COMERCIO CON DESPLEGABLES
c_rank1, c_rank2 = st.columns([2, 1])
with c_rank1:
    st.subheader(txt.get("seccion_ranking_titulo", "🏪 Rendimiento y Semáforos por Comercio"))
with c_rank2:
    criterio_orden = st.selectbox(
        "Ordenar ranking por:",
        [
            "💰 Mayor Colocador ($ Colocado)",
            "💼 Mayor Venta Total ($ Pipeline)",
            "🎯 Mayor Approval Rate (% Aprobación)",
            "⚡ Mayor Conversión (%)",
            "📋 Mayor Cantidad de Consultas"
        ],
        index=0
    )

# Colocación anual y del mes actual: independiente del filtro de fechas de alta.
import calendar
fecha_avance = datetime.datetime.now(ZoneInfo("America/Argentina/Buenos_Aires")).date()
fraccion_mes = fecha_avance.day / calendar.monthrange(fecha_avance.year, fecha_avance.month)[1]
ruta_historial = os.path.join(os.path.dirname(__file__), "downloads", "HistorialLiquidaciones.xlsx")
historial = cargar_historial_liquidaciones(os.stat(ruta_historial).st_mtime_ns if os.path.exists(ruta_historial) else None)
fecha_inicio_historial = pd.to_datetime(historial.get("FechaLiquidacion", pd.Series(dtype=str)), format="mixed", dayfirst=True, errors="coerce").min()
historial_anual_disponible = pd.notna(fecha_inicio_historial) and fecha_inicio_historial.date() <= datetime.date(fecha_avance.year, 1, 31)
base_colocacion = pd.concat([historial, df_base], ignore_index=True).drop_duplicates("Solicitud", keep="last") if not historial.empty else df_base.copy()
fechas_liquidacion = pd.to_datetime(base_colocacion["FechaLiquidacion"], format="mixed", dayfirst=True, errors="coerce")
base_colocacion["FechaLiquidacion"] = fechas_liquidacion
base_colocacion = base_colocacion[
    base_colocacion["Estado"].map(normalizar_estado).eq("LIQUIDADA")
    & fechas_liquidacion.dt.year.eq(fecha_avance.year)
    & fechas_liquidacion.dt.date.le(fecha_avance)
].copy()
base_colocacion["MesLiquidacion"] = pd.to_datetime(base_colocacion["FechaLiquidacion"]).dt.month
anual_comercio = base_colocacion.groupby("Comercio")["Monto"].sum()
mes_comercio = base_colocacion[base_colocacion["MesLiquidacion"].eq(fecha_avance.month)].groupby("Comercio")["Monto"].sum()

def avance_colocacion(acumulado_anual, acumulado_mes):
    promedio = float(acumulado_anual) / fecha_avance.month if historial_anual_disponible else None
    esperado = promedio * fraccion_mes if promedio is not None else 0
    avance = float(acumulado_mes) / esperado * 100 if esperado > 0 else None
    return {
        "Promedio mensual ($)": promedio,
        "Colocado este mes ($)": float(acumulado_mes),
        "Avance esperado (%)": avance,
        "Semáforo colocación": ("⚪ Sin historial anual" if not historial_anual_disponible else "⚪ Sin base") if avance is None else ("🟢 ↑" if avance >= 100 else "🟡 ↓" if avance >= 90 else "🔴 ↓"),
    }

fechas_actividad = pd.to_datetime(base_diaria["FechaAlta"], dayfirst=True, errors="coerce").dt.date
fecha_ayer = hoy_argentina - datetime.timedelta(days=1)
consultas_dia_por_comercio = base_diaria[fechas_actividad.eq(hoy_argentina)].groupby("Comercio").size()
consultas_ayer_por_comercio = base_diaria[fechas_actividad.eq(fecha_ayer)].groupby("Comercio").size()
fechas_disponibles = pd.to_datetime(df_raw["FechaAlta"], dayfirst=True, errors="coerce").dropna()
actividad_comparable = not fechas_disponibles.empty and fechas_disponibles.min().date() <= fecha_ayer and reporte_actualizado.date() >= hoy_argentina

def indicador_actividad(actual, anterior, comparable=True):
    if not comparable:
        return "", "#1e293b"
    if actual > anterior:
        return "↑", "#15803d"
    if actual < anterior:
        return "↓", "#dc2626"
    return "—", "#a16207"


def rango_anterior_consultas(inicio, fin):
    if inicio == fin:
        return inicio - datetime.timedelta(days=1), inicio - datetime.timedelta(days=1)
    if inicio.month == 1 and inicio.day == 1 and fin.year == inicio.year and fin.month > 1:
        return datetime.date(inicio.year - 1, 1, 1), datetime.date(inicio.year - 1, 12, 31)
    if inicio.day == 1 and (inicio.year, inicio.month) == (fin.year, fin.month):
        ultimo = inicio - datetime.timedelta(days=1)
        return ultimo.replace(day=1), ultimo
    anterior_hasta = inicio - datetime.timedelta(days=1)
    return anterior_hasta - (fin - inicio), anterior_hasta

anterior_desde, anterior_hasta = rango_anterior_consultas(f_desde, f_hasta)
consultas_periodo_anterior = base_diaria[fechas_actividad.ge(anterior_desde) & fechas_actividad.le(anterior_hasta)].groupby("Comercio").size()
periodo_comparable = not fechas_disponibles.empty and fechas_disponibles.min().date() <= anterior_desde and reporte_actualizado.date() >= anterior_hasta
resumen_comercios = []
for com_nombre, group in df_filtrado.groupby("Comercio"):
    sub_rec = [{"estado": r["Estado"], "monto": r["Monto"]} for _, r in group.iterrows()]
    sub_m = calcular_metricas(sub_rec)
    loc_val = group["Localidad"].iloc[0] if "Localidad" in group.columns else "-"
    resumen_comercios.append({
        "Comercio": com_nombre,
        "Localidad": loc_val,
        "Consultas del día": int(consultas_dia_por_comercio.get(com_nombre, 0)),
        "Consultas totales": sub_m["total_consultas"],
        "Consultas con oferta": sub_m["total_con_oferta"],
        "Rechazadas": sub_m["rechazadas"],
        "Incompletas": sub_m["incompletas"],
        "En Análisis": sub_m["en_analisis"],
        "Aprobadas": sub_m["aprobadas"],
        "Liquidadas": sub_m["liquidadas"],
        "Venta Total ($)": sub_m["monto_venta_total"],
        "Total Colocado ($)": sub_m["monto_liquidado"],
        **avance_colocacion(anual_comercio.get(com_nombre, 0), mes_comercio.get(com_nombre, 0)),
        "Ticket promedio ($)": ticket_promedio(sub_m["monto_venta_total"], sub_m["venta_total_cant"]),
        "Aprobación (%)": sub_m["pct_aprobacion"],
        "Conversión (%)": sub_m["pct_conversion_bruta"],
        "Semáforo": calcular_semaforo(sub_m["pct_conversion_bruta"])["icono"]
    })

df_resumen = pd.DataFrame(resumen_comercios)

# Aplicar orden según criterio seleccionado
if "Mayor Colocador" in criterio_orden:
    df_resumen = df_resumen.sort_values("Total Colocado ($)", ascending=False)
elif "Mayor Venta Total" in criterio_orden:
    df_resumen = df_resumen.sort_values("Venta Total ($)", ascending=False)
elif "Approval Rate" in criterio_orden:
    df_resumen = df_resumen.sort_values("Aprobación (%)", ascending=False)
elif "Mayor Conversión" in criterio_orden:
    df_resumen = df_resumen.sort_values("Conversión (%)", ascending=False)
elif "Cantidad de Consultas" in criterio_orden:
    df_resumen = df_resumen.sort_values("Consultas totales", ascending=False)
else:
    df_resumen = df_resumen.sort_values("Total Colocado ($)", ascending=False)

# Fila consolidada de totales al pie de la tabla
if not df_resumen.empty:
    tot_consultas = int(df_resumen["Consultas totales"].sum())
    tot_rechazadas = int(df_resumen["Rechazadas"].sum())
    tot_incompletas = int(df_resumen["Incompletas"].sum())
    tot_en_analisis = int(df_resumen["En Análisis"].sum())
    tot_aprobadas = int(df_resumen["Aprobadas"].sum())
    tot_liquidadas = int(df_resumen["Liquidadas"].sum())
    tot_venta_total = float(df_resumen["Venta Total ($)"].sum())
    tot_colocado = float(df_resumen["Total Colocado ($)"].sum())
    
    tot_con_oferta = tot_incompletas + tot_en_analisis + tot_aprobadas + tot_liquidadas
    pct_aprob_global = (tot_con_oferta / tot_consultas * 100) if tot_consultas > 0 else 0.0
    pct_conv_global = (tot_liquidadas / tot_consultas * 100) if tot_consultas > 0 else 0.0
    
    fila_total = pd.DataFrame([{
        "Comercio": "📌 TOTAL CONSOLIDADO",
        "Localidad": "Toda la Cartera",
        "Consultas del día": int(df_resumen["Consultas del día"].sum()),
        "Consultas totales": tot_consultas,
        "Consultas con oferta": tot_con_oferta,
        "Rechazadas": tot_rechazadas,
        "Incompletas": tot_incompletas,
        "En Análisis": tot_en_analisis,
        "Aprobadas": tot_aprobadas,
        "Liquidadas": tot_liquidadas,
        "Venta Total ($)": tot_venta_total,
        "Total Colocado ($)": tot_colocado,
        **avance_colocacion(anual_comercio.reindex(df_resumen["Comercio"], fill_value=0).sum(), mes_comercio.reindex(df_resumen["Comercio"], fill_value=0).sum()),
        "Ticket promedio ($)": ticket_promedio(tot_venta_total, tot_en_analisis + tot_aprobadas + tot_liquidadas),
        "Aprobación (%)": round(pct_aprob_global, 1),
        "Conversión (%)": round(pct_conv_global, 1),
        "Semáforo": calcular_semaforo(pct_conv_global)["icono"]
    }])
    
    df_resumen_con_total = pd.concat([df_resumen, fila_total], ignore_index=True)
else:
    df_resumen_con_total = df_resumen


# Mostrar tabla resumen
with st.expander("Más info", expanded=False):
    if not historial_anual_disponible:
        st.info("Falta el historial desde enero. Actualizá los datos del portal para calcular el promedio y el avance; no se consideran meses faltantes como meses sin ventas.")
    st.caption(f"Colocación mensual al {fecha_avance:%d/%m/%Y}: liquidados del año ÷ {fecha_avance.month} meses (incluye meses sin colocación). Esperado hoy: {fraccion_mes:.1%} del promedio. Avance: 🟢 ≥100% · 🟡 ≥90% y <100% · 🔴 <90% · ⚪ sin base. Usa fecha de liquidación y el historial disponible, independientemente del rango de fechas de alta.")
    st.caption(f"Consultas del día: parcial de hoy {hoy_argentina:%d/%m/%Y}, independiente del filtro de fechas. Consultas conserva el total del período elegido. Actividad: ↑ verde si supera ayer, — amarillo si iguala, ↓ rojo si baja. Compara hoy parcial con ayer completo según fecha de alta; la flecha aparece cuando hay datos de ambos días.")
    st.caption(f"Actividad del período: {f_desde:%d/%m/%Y}–{f_hasta:%d/%m/%Y} contra {anterior_desde:%d/%m/%Y}–{anterior_hasta:%d/%m/%Y}. Un mes se compara con el mes anterior; un año, con el anterior; otros rangos, con igual cantidad de días previos. Se cuentan todos los estados.")
st.caption("Tocá una fila para ver sus créditos liquidados, en análisis y aprobados.")
estilo_resumen = df_resumen_con_total.style.format({
        "Venta Total ($)": pesos,
        "Total Colocado ($)": pesos,
        "Promedio mensual ($)": pesos,
        "Colocado este mes ($)": pesos,
        "Avance esperado (%)": lambda valor: "—" if pd.isna(valor) else f"{valor:.1f}%",
        "Ticket promedio ($)": pesos,
        "Aprobación (%)": "{:.1f}%",
        "Conversión (%)": "{:.1f}%",
    })
for indice, fila in df_resumen_con_total.iterrows():
    anterior = int(consultas_ayer_por_comercio.reindex(df_resumen["Comercio"], fill_value=0).sum()) if indice == len(df_resumen) else int(consultas_ayer_por_comercio.get(fila["Comercio"], 0))
    simbolo, color = indicador_actividad(int(fila["Consultas del día"]), anterior, actividad_comparable)
    celda = pd.IndexSlice[[indice], ["Consultas del día"]]
    estilo_resumen = estilo_resumen.format(lambda valor, flecha=simbolo: f"{int(valor):,} {flecha}".strip(), subset=celda).set_properties(subset=celda, color=color, **{"font-weight": "bold"})
    anterior_periodo = int(consultas_periodo_anterior.reindex(df_resumen["Comercio"], fill_value=0).sum()) if indice == len(df_resumen) else int(consultas_periodo_anterior.get(fila["Comercio"], 0))
    simbolo_periodo, color_periodo = indicador_actividad(int(fila["Consultas totales"]), anterior_periodo, periodo_comparable)
    celda_periodo = pd.IndexSlice[[indice], ["Consultas totales"]]
    estilo_resumen = estilo_resumen.format(lambda valor, flecha=simbolo_periodo: f"{int(valor):,} {flecha}".strip(), subset=celda_periodo).set_properties(subset=celda_periodo, color=color_periodo, **{"font-weight": "bold"})
seleccion_comercio = st.dataframe(
    estilo_resumen,
    use_container_width=True,
    height=340,
    hide_index=True,
    column_order=["Comercio", "Consultas totales", "Consultas del día", "Consultas con oferta", "Aprobadas", "En Análisis", "Liquidadas", "Rechazadas", "Incompletas", "Aprobación (%)", "Conversión (%)", "Total Colocado ($)", "Venta Total ($)", "Ticket promedio ($)", "Promedio mensual ($)", "Colocado este mes ($)", "Avance esperado (%)", "Semáforo colocación", "Semáforo", "Localidad"],
    on_select="rerun",
    selection_mode="single-cell",
    key="ranking_" + str(hash((criterio_orden, vista_seleccionada, str(f_desde), str(f_hasta), tuple(df_resumen_con_total.get("Comercio", []))))),
)

descargar_excel(df_resumen_con_total, "rendimiento_comercios", "excel_ranking")

filas_seleccionadas = [celda[0] for celda in seleccion_comercio.selection.cells]
if filas_seleccionadas:
    fila = filas_seleccionadas[0]
    if fila < len(df_resumen_con_total):
        comercio_elegido = df_resumen_con_total.iloc[fila]["Comercio"]
        es_total = fila >= len(df_resumen)
        detalle = df_filtrado.copy() if es_total else df_filtrado[df_filtrado["Comercio"] == comercio_elegido].copy()
        detalle["Estado_normalizado"] = detalle["Estado"].map(normalizar_estado)
        detalle = detalle[detalle["Estado_normalizado"].isin(["LIQUIDADA", "EN_ANALISIS", "APROBADA"])]
        st.subheader("Créditos — " + ("Total consolidado" if es_total else str(comercio_elegido)))
        st.caption("Detalle del período seleccionado. Solo liquidados, en análisis y aprobados.")
        if detalle.empty:
            st.info("Este comercio no tiene créditos en esos estados dentro del período seleccionado.")
        else:
            columnas_estado = st.columns(3)
            for columna, estado, titulo in zip(columnas_estado, ["LIQUIDADA", "EN_ANALISIS", "APROBADA"], ["Liquidados", "En análisis", "Aprobados"]):
                subset = detalle[detalle["Estado_normalizado"] == estado]
                columna.metric(titulo, str(len(subset)), f"{pesos(subset['Monto'].sum())}", delta_color="off")
            st.metric("Ticket promedio del comercio", pesos(detalle["Monto"].mean()))
            detalle = detalle.sort_values("FechaAlta", ascending=False)
            campos = ["Solicitud", "Nombre", "Comercio", "Estado", "Monto", "FechaAlta", "FechaAprobacion", "FechaLiquidacion"]
            detalle = detalle[[c for c in campos if c in detalle.columns]].copy()
            for c in ["FechaAlta", "FechaAprobacion", "FechaLiquidacion"]:
                if c in detalle:
                    detalle[c] = pd.to_datetime(detalle[c], dayfirst=True, errors="coerce").dt.strftime("%d/%m/%Y").fillna("—")
            detalle["Nombre"] = detalle["Nombre"].fillna("No informado")
            detalle = detalle.rename(columns={"Nombre": "Nombre y apellido del cliente", "Solicitud": "Solicitud Nº", "FechaAlta": "Fecha de alta", "FechaAprobacion": "Fecha de aprobación", "FechaLiquidacion": "Fecha de liquidación"})
            st.dataframe(detalle.style.format({"Monto": pesos}), hide_index=True, use_container_width=True)
            st.markdown(f"**Total del monto colocado: {pesos(detalle['Monto'].sum())}**")
            descargar_excel(detalle, "creditos_comercio", "excel_creditos_comercio")

# Resumen visual destacado justo abajo de la tabla
if not solo_general:
    st.markdown(f"""
    <div style="background-color:#f8fafc; border:1px solid #cbd5e1; border-left:5px solid #059669; border-radius:8px; padding:12px 18px; margin-top:8px; margin-bottom:20px;">
        <div style="font-size:0.9rem; color:#64748b; font-weight:600; text-transform:uppercase;">📊 Consolidado de Colocación (Mi Cartera)</div>
        <div style="display:flex; gap:20px; flex-wrap:wrap; margin-top:6px;">
            <div><span style="color:#64748b; font-size:0.85rem;">Total Colocado (Liquidado):</span><br><b style="font-size:1.3rem; color:#059669;">{pesos(m_cartera_full['monto_liquidado'])}</b> <span style="font-size:0.85rem; color:#475569;">({m_cartera_full['total_liquidado_cant']} créditos)</span></div>
            <div><span style="color:#64748b; font-size:0.85rem;">Venta Total en Curso (Pipeline):</span><br><b style="font-size:1.3rem; color:#1e293b;">{pesos(m_cartera_full['monto_venta_total'])}</b> <span style="font-size:0.85rem; color:#475569;">({m_cartera_full['venta_total_cant']} operaciones)</span></div>
            <div><span style="color:#64748b; font-size:0.85rem;">Consultas Totales Cartera:</span><br><b style="font-size:1.3rem; color:#1e293b;">{m_cartera_full['total_consultas']:,}</b></div>
        </div>
    </div>
    """, unsafe_allow_html=True)



# SECCIÓN DE DESPLEGABLES INDIVIDUALES POR COMERCIO
st.subheader("🔍 Fichas Desplegables de Rendimiento Mes a Mes (MoM)")
for _, row in df_resumen.head(10).iterrows():
    c_name = row["Comercio"]
    with st.expander(f"{row['Semáforo']} **{c_name}** | Venta: {pesos(row['Venta Total ($)'])} | Conv: {row['Conversión (%)']:.1f}% | Consultas: {row['Consultas totales']}"):
        col_a, col_b, col_c, col_d = st.columns(4)
        col_a.metric("Consultas totales", f"{row['Consultas totales']}", "Base del mes")
        col_b.metric("Venta Total", f"{pesos(row['Venta Total ($)'])}", f"{row['Liquidadas']} liquidadas")
        col_c.metric("Approval Rate", f"{row['Aprobación (%)']:.1f}%", "Tasa de Oferta")
        col_d.metric("Conversión", f"{row['Conversión (%)']:.1f}%", row["Semáforo"])
        
        st.caption("Estados detallados de solicitudes:")
        st.write({
            "Rechazadas": row["Rechazadas"],
            "Incompletas (con oferta)": row["Incompletas"],
            "En análisis": row["En Análisis"],
            "Aprobadas": row["Aprobadas"],
            "Liquidadas": row["Liquidadas"]
        })
